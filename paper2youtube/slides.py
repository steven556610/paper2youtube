"""Stage 4 — render slides_plan → PPTX + per-slide PNG (PIL)."""

from __future__ import annotations

import re
from pathlib import Path

from .script_builder import SlidesPlan

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:  # pragma: no cover
    Image = ImageDraw = ImageFont = None  # type: ignore


SLIDE_W, SLIDE_H = 1280, 720
BG = (255, 255, 255)
FG = (24, 24, 24)
ACCENT = (153, 0, 0)
MUTED = (90, 90, 90)


def render_slides(plan: SlidesPlan, workdir: Path, *, start_index: int = 1) -> list[Path]:
    """Write slide_NN.png under workdir/slides/ and output_slides.pptx at root."""
    workdir = Path(workdir)
    slides_dir = workdir / "slides"
    slides_dir.mkdir(exist_ok=True)

    pngs: list[Path] = []
    for i, item in enumerate(plan.items, start=start_index):
        png = slides_dir / f"slide_{i:02d}.png"
        render_slide_png(item.slide_markdown, item.title, png, figure_path=item.figure_path, workdir=workdir)
        pngs.append(png)

    pptx_path = workdir / "output_slides.pptx"
    render_pptx(plan, pptx_path, workdir=workdir)
    return pngs


def render_slide_png(
    markdown: str,
    title: str,
    out_path: Path,
    *,
    figure_path: str | None = None,
    workdir: Path | None = None,
) -> Path:
    if Image is None:  # pragma: no cover
        raise RuntimeError("Pillow is required to render slide PNGs")

    img = Image.new("RGB", (SLIDE_W, SLIDE_H), BG)
    draw = ImageDraw.Draw(img)
    font_title = _font(42)
    font_body = _font(26)
    font_small = _font(20)

    # header bar
    draw.rectangle([0, 0, SLIDE_W, 8], fill=ACCENT)
    _wrap_draw(draw, title, 60, 36, font_title, FG, max_width=SLIDE_W - 120)
    draw.line([(60, 100), (SLIDE_W - 60, 100)], fill=ACCENT, width=2)

    body_y = 130
    bullets = _md_bullets(markdown)
    fig_path = _resolve_figure(markdown, figure_path, workdir)

    if fig_path and Path(fig_path).exists():
        try:
            fig = Image.open(fig_path).convert("RGB")
            fig.thumbnail((520, 360))
            x = SLIDE_W - fig.width - 50
            y = (SLIDE_H - fig.height) // 2 + 20
            img.paste(fig, (x, y))
            draw.rectangle([x - 4, y - 4, x + fig.width + 3, y + fig.height + 3], outline=MUTED, width=2)
        except Exception:
            fig_path = None

    text_width = 520 if fig_path else SLIDE_W - 120
    for b in bullets[:6]:
        body_y = _wrap_draw(draw, f"• {b}", 60, body_y, font_body, FG, max_width=text_width, line_gap=8)
        body_y += 10

    # caption strip
    draw.rectangle([0, SLIDE_H - 42, SLIDE_W, SLIDE_H], fill=(245, 245, 245))
    _wrap_draw(draw, "paper2youtube", 20, SLIDE_H - 34, font_small, MUTED, max_width=200)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    img.save(out_path, format="PNG")
    return out_path


def render_pptx(plan: SlidesPlan, out_path: Path, *, workdir: Path | None = None) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]  # blank

    for item in plan.items:
        slide = prs.slides.add_slide(blank)
        _add_text(slide, item.title, left=0.6, top=0.3, width=12.1, height=0.9, size=28, bold=True)
        bullets = _md_bullets(item.slide_markdown)
        body = "\n".join(f"• {b}" for b in bullets[:7])
        _add_text(slide, body, left=0.6, top=1.3, width=7.2 if item.figure_path else 12.1, height=5.5, size=16)
        fig = _resolve_figure(item.slide_markdown, item.figure_path, workdir)
        if fig and Path(fig).exists():
            try:
                slide.shapes.add_picture(str(fig), Inches(8.2), Inches(1.6), width=Inches(4.5))
            except Exception:
                pass

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


def _add_text(slide, text: str, *, left: float, top: float, width: float, height: float, size: int, bold: bool = False):
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold


def _md_bullets(markdown: str) -> list[str]:
    out = []
    for line in (markdown or "").splitlines():
        s = line.strip()
        if s.startswith("#"):
            continue
        if s.startswith("!"):
            continue
        if s.startswith(">"):
            s = s.lstrip("> ").strip()
        s = re.sub(r"^\s*[-*+]\s+", "", s)
        s = re.sub(r"\*\*([^*]+)\*\*", r"\1", s)
        s = re.sub(r"\*([^*]+)\*", r"\1", s)
        s = s.strip()
        if not s or s.startswith("```"):
            continue
        if s.lower().startswith("figures/"):
            continue
        out.append(s[:220])
    return out


def _resolve_figure(markdown: str, figure_path: str | None, workdir: Path | None) -> str | None:
    if figure_path and Path(figure_path).exists():
        return figure_path
    m = re.search(r"!\[[^\]]*\]\(([^)]+)\)", markdown or "")
    if not m:
        return None
    rel = m.group(1)
    if workdir:
        cand = workdir / rel
        if cand.exists():
            return str(cand)
        cand = workdir / "figures" / Path(rel).name
        if cand.exists():
            return str(cand)
    return rel if Path(rel).exists() else None


def _font(size: int):
    if ImageFont is None:  # pragma: no cover
        return None
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/System/Library/Fonts/PingFang.ttc",
    ]
    for c in candidates:
        if Path(c).exists():
            try:
                return ImageFont.truetype(c, size)
            except Exception:
                continue
    return ImageFont.load_default()


def _wrap_draw(draw, text: str, x: int, y: int, font, fill, *, max_width: int, line_gap: int = 4) -> int:
    words = list(text)
    if not words:
        return y
    # CJK-aware wrap: break by characters when needed
    line = ""
    for ch in text:
        trial = line + ch
        w = draw.textlength(trial, font=font) if hasattr(draw, "textlength") else font.getlength(trial)
        if w > max_width and line:
            draw.text((x, y), line, font=font, fill=fill)
            y += getattr(font, "size", 22) + line_gap
            line = ch
        else:
            line = trial
    if line:
        draw.text((x, y), line, font=font, fill=fill)
        y += getattr(font, "size", 22) + line_gap
    return y